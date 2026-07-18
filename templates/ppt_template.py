"""Presentation theme and slide layout helpers for python-pptx.

This module keeps styling concerns separate from the exporter so that the
theme can evolve without touching export logic.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Sequence

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class PresentationTheme:
    """Color and font theme."""

    title_color: RGBColor = RGBColor(0x1A, 0x47, 0x5F)  # dark teal
    text_color: RGBColor = RGBColor(0x33, 0x33, 0x33)  # near black
    accent_color: RGBColor = RGBColor(0x4A, 0x90, 0xA4)  # teal accent
    background_color: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)
    font_name: str = "Helvetica Neue"
    title_font_size: int = 32
    subtitle_font_size: int = 18
    body_font_size: int = 20
    footer_font_size: int = 12


DEFAULT_THEME = PresentationTheme()


def _set_text_frame_style(
    text_frame,
    font_name: str,
    font_size: int,
    color: RGBColor,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    text_frame.word_wrap = True
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.color.rgb = color


def apply_title_slide(slide: Slide, title: str, subtitle: str, theme: PresentationTheme = DEFAULT_THEME) -> None:
    """Style a title slide."""
    # Clear default placeholders and add custom text boxes for reliability.
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.placeholder_format is not None:
            sp = shape
            sp.text = ""

    left = Inches(0.75)
    top = Inches(2.0)
    width = Inches(8.5)
    height = Inches(2.0)
    title_box = slide.shapes.add_textbox(left, top, width, height)
    tf = title_box.text_frame
    tf.text = title
    _set_text_frame_style(
        tf,
        theme.font_name,
        theme.title_font_size,
        theme.title_color,
        PP_ALIGN.LEFT,
    )

    subtitle_top = Inches(4.25)
    sub_box = slide.shapes.add_textbox(left, subtitle_top, width, Inches(1.5))
    stf = sub_box.text_frame
    stf.text = subtitle
    _set_text_frame_style(
        stf,
        theme.font_name,
        theme.subtitle_font_size,
        theme.text_color,
        PP_ALIGN.LEFT,
    )


def apply_content_slide(
    slide: Slide,
    title: str,
    bullets: Sequence[str],
    footer: str,
    theme: PresentationTheme = DEFAULT_THEME,
) -> None:
    """Style a standard content slide."""
    left = Inches(0.6)
    top = Inches(0.5)
    width = Inches(8.8)
    height = Inches(1.0)
    title_box = slide.shapes.add_textbox(left, top, width, height)
    tf = title_box.text_frame
    tf.text = title
    _set_text_frame_style(
        tf,
        theme.font_name,
        theme.title_font_size,
        theme.title_color,
        PP_ALIGN.LEFT,
    )

    body_top = Inches(1.5)
    body_height = Inches(5.0)
    body_box = slide.shapes.add_textbox(left, body_top, width, body_height)
    btf = body_box.text_frame
    btf.text = ""
    for idx, bullet in enumerate(bullets):
        p = btf.paragraphs[idx] if idx < len(btf.paragraphs) else btf.add_paragraph()
        p.text = f"• {bullet}"
        p.level = 0
        p.font.name = theme.font_name
        p.font.size = Pt(theme.body_font_size)
        p.font.color.rgb = theme.text_color

    if footer:
        footer_top = Inches(6.8)
        footer_box = slide.shapes.add_textbox(left, footer_top, width, Inches(0.4))
        ftf = footer_box.text_frame
        ftf.text = footer
        _set_text_frame_style(
            ftf,
            theme.font_name,
            theme.footer_font_size,
            theme.accent_color,
            PP_ALIGN.LEFT,
        )
