"""Negative and closure tests for the strict PubMed evidence contract."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import httpx
from pptx import Presentation
import pytest

from src.pipeline.generator import LessonGenerator
from src.providers.base import ProviderError
from src.pubmed.briefs import generate_evidence_briefs
from src.pubmed.builder import PubMedLessonBuilder
from src.pubmed.client import (
    PubMedClient,
    PubMedClientConfig,
    PubMedError,
    PubMedRetrievalResult,
)
from src.pubmed.models import PubMedSearchQuery
from src.pubmed.normalizer import (
    infer_evidence_level,
    infer_study_type,
    parse_efetch_xml,
)
from src.pubmed.traceability import validate_pubmed_traceability
from src.schemas.lesson import Language, LessonPackage


FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures"
EFETCH_XML = (FIXTURE_DIR / "pubmed_efetch.xml").read_text(encoding="utf-8")
RECORDS = parse_efetch_xml(EFETCH_XML)


class _FakeResponse:
    def __init__(self, payload: Any) -> None:
        self.status_code = 200
        self.text = ""
        self._payload = payload

    def json(self) -> Any:
        return self._payload


class _SingleResponseClient:
    def __init__(self, payload: Any) -> None:
        self.payload = payload
        self.calls = 0

    def get(self, url: str, **kwargs: Any) -> _FakeResponse:
        self.calls += 1
        return _FakeResponse(self.payload)


class _ConnectFailureClient:
    def __init__(self) -> None:
        self.calls = 0

    def get(self, url: str, **kwargs: Any) -> _FakeResponse:
        self.calls += 1
        raise httpx.ConnectError(
            "connection refused",
            request=httpx.Request("GET", url),
        )


def _client_with_fake_http(fake: Any) -> PubMedClient:
    client = PubMedClient(PubMedClientConfig(api_key="fixture"))
    client._client = fake  # type: ignore[assignment]
    client._rate_limit = lambda: None
    return client


def _build_package() -> LessonPackage:
    return PubMedLessonBuilder(
        topic="Dietary Fiber and Gut Health",
        audience="general adults",
        language=Language.EN,
        duration_minutes=10,
        records=RECORDS,
    ).build()


def _fixture_result(
    date_range: tuple[str | None, str | None] = (None, None),
) -> PubMedRetrievalResult:
    return PubMedRetrievalResult(
        query=PubMedSearchQuery(
            topic="Dietary Fiber and Gut Health",
            final_query="fixture query",
            date_range=date_range,
            requested_max_count=3,
        ),
        raw_esearch={"esearchresult": {"idlist": [record.pmid for record in RECORDS]}},
        raw_efetch=EFETCH_XML,
        records=RECORDS,
        briefs=generate_evidence_briefs(RECORDS, "Dietary Fiber and Gut Health"),
    )


def _fake_pubmed_client_class() -> Any:
    class FakeClient:
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            pass

        def __enter__(self) -> "FakeClient":
            return self

        def __exit__(self, *args: Any, **kwargs: Any) -> None:
            pass

        def retrieve(self, *args: Any, **kwargs: Any) -> PubMedRetrievalResult:
            return _fixture_result(kwargs.get("date_range", (None, None)))

    return FakeClient


def test_esearch_requires_well_formed_schema_and_numeric_pmids():
    malformed_payloads = [
        {},
        {"esearchresult": {}},
        {"esearchresult": {"idlist": "123"}},
        {"esearchresult": {"idlist": ["123", "bad-pmid"]}},
    ]
    for payload in malformed_payloads:
        client = _client_with_fake_http(_SingleResponseClient(payload))
        with pytest.raises(PubMedError, match="Malformed ESearch response"):
            client.search("fixture", max_results=3)


def test_esearch_deduplicates_pmids_preserving_order():
    fake = _SingleResponseClient(
        {"esearchresult": {"idlist": ["123", "123", "456"]}}
    )
    client = _client_with_fake_http(fake)
    _, pmids = client.search("fixture", max_results=3)
    assert pmids == ["123", "456"]


def test_connect_error_is_descriptive_and_not_retried():
    fake = _ConnectFailureClient()
    client = _client_with_fake_http(fake)
    with pytest.raises(PubMedError, match="Non-retryable NCBI network error"):
        client.search("fixture", max_results=1)
    assert fake.calls == 1


def test_invalid_date_range_rejected_before_network():
    client = PubMedClient(PubMedClientConfig(api_key="fixture"))
    with pytest.raises(PubMedError, match="date range"):
        client.build_query("fiber", date_range=("2025/13/01", None))
    with pytest.raises(PubMedError, match="date range"):
        client.build_query("fiber", date_range=("2025", "2024"))


def test_non_randomized_title_is_not_misclassified_as_rct():
    title = "A non-randomized nutrition intervention"
    assert infer_study_type([], title) == "not_reported"
    assert infer_evidence_level([], title) == "not_reported"


def test_missing_title_stays_null_and_invalid_pmid_fails():
    missing_title_xml = """
    <PubmedArticleSet><PubmedArticle><MedlineCitation>
      <PMID>12345</PMID><Article><Language>eng</Language></Article>
    </MedlineCitation></PubmedArticle></PubmedArticleSet>
    """
    records = parse_efetch_xml(missing_title_xml)
    assert len(records) == 1
    assert records[0].title is None

    invalid_pmid_xml = missing_title_xml.replace("12345", "abc")
    with pytest.raises(ValueError, match="Invalid PMID"):
        parse_efetch_xml(invalid_pmid_xml)


@pytest.mark.parametrize(
    "case",
    [
        "claim_without_citation",
        "unknown_citation",
        "reference_without_pmid",
        "reference_metadata_mismatch",
        "bad_slide_mapping",
    ],
)
def test_strict_traceability_rejects_integrity_breaks(case: str):
    package = _build_package()
    validate_pubmed_traceability(package, RECORDS)

    if case == "claim_without_citation":
        package.claims[0].citation_ids = []
    elif case == "unknown_citation":
        package.claims[0].citation_ids = ["PMID_99999999"]
    elif case == "reference_without_pmid":
        package.references[0].pmid = None
    elif case == "reference_metadata_mismatch":
        package.references[0].title = "Mismatched title"
    elif case == "bad_slide_mapping":
        package.claims[0].slide_numbers = [999]

    with pytest.raises(ValueError, match="PubMed traceability validation failed"):
        validate_pubmed_traceability(package, RECORDS)


def test_claim_slide_mapping_is_visible_and_excludes_reference_slide():
    package = _build_package()
    reference_slide_numbers = {
        slide.slide_number for slide in package.slides if slide.slide_type == "references"
    }
    for claim in package.claims:
        assert claim.slide_numbers
        assert not (set(claim.slide_numbers) & reference_slide_numbers)
        for slide_number in claim.slide_numbers:
            slide = next(
                slide for slide in package.slides if slide.slide_number == slide_number
            )
            assert claim.claim_text in "\n".join(slide.key_points)
            assert set(claim.citation_ids) & set(slide.citation_ids)


def test_pubmed_export_always_writes_canonical_request_and_ppt_citations(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
):
    monkeypatch.setattr(
        "src.pipeline.generator.PubMedClient",
        _fake_pubmed_client_class(),
    )
    generator = LessonGenerator(output_dir=tmp_path)
    package = generator.generate(
        topic="Dietary Fiber and Gut Health",
        audience="general adults",
        language=Language.EN,
        duration_minutes=10,
        evidence_source="pubmed",
        max_pubmed_results=3,
        pubmed_date_range=("2010", "2025"),
        pubmed_timeout=12.5,
    )
    out = generator.export_package(package)

    request = json.loads((out / "request.json").read_text(encoding="utf-8"))
    assert request["evidence_source"] == "pubmed"
    assert request["date_from"] == "2010"
    assert request["date_to"] == "2025"
    assert request["pubmed_timeout_seconds"] == 12.5

    presentation = Presentation(out / "nutrition_lesson.pptx")
    ppt_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    for record in RECORDS:
        assert f"PMID: {record.pmid}" in ppt_text
    for claim in package.claims:
        assert claim.claim_text in ppt_text


def test_detached_pubmed_package_cannot_silently_omit_retrieval_tree(tmp_path: Path):
    package = _build_package()
    detached = LessonPackage.model_validate(package.model_dump(mode="json"))
    generator = LessonGenerator(output_dir=tmp_path)
    with pytest.raises(ProviderError, match="original validated retrieval result"):
        generator.export_package(detached)
