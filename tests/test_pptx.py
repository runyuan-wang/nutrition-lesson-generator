"""Tests for PowerPoint export and OOXML validation."""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation

from src.exporters.ppt import PptxExporter
from src.providers.mock_provider import MockProvider
from src.schemas.lesson import Language


def test_pptx_export_and_reopen(tmp_path: Path):
    provider = MockProvider()
    package = provider.generate(
        topic="Dietary Fiber and Gut Health",
        audience="general adults",
        language=Language.EN,
    )
    out = tmp_path / "lesson.pptx"
    PptxExporter().export(package, out)

    assert out.exists()
    assert out.stat().st_size > 0

    prs = Presentation(str(out))
    assert len(prs.slides) == len(package.slides)

    first_slide = prs.slides[0]
    assert first_slide.shapes


def test_pptx_is_valid_ooxml_zip(tmp_path: Path):
    provider = MockProvider()
    package = provider.generate(
        topic="Dietary Fiber and Gut Health",
        audience="general adults",
        language=Language.EN,
    )
    out = tmp_path / "lesson.pptx"
    PptxExporter().export(package, out)

    with zipfile.ZipFile(out, "r") as zf:
        names = zf.namelist()
        assert "[Content_Types].xml" in names
        assert "ppt/presentation.xml" in names
        slide_parts = [n for n in names if n.startswith("ppt/slides/slide")]
        assert len(slide_parts) == len(package.slides)


def test_pptx_has_references_slide(tmp_path: Path):
    provider = MockProvider()
    package = provider.generate(
        topic="Dietary Fiber and Gut Health",
        audience="general adults",
        language=Language.EN,
    )
    out = tmp_path / "lesson.pptx"
    PptxExporter().export(package, out)

    prs = Presentation(str(out))
    titles = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                titles.append(shape.text_frame.text.strip().split("\n")[0])
    assert any("Reference" in t for t in titles)
