"""PowerPoint exporter using python-pptx."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from src.schemas.lesson import CitationRecord, LessonPackage, LessonSlide
from templates.ppt_template import (
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    apply_content_slide,
    apply_evidence_slide,
    apply_objectives_slide,
    apply_references_slide,
    apply_summary_slide,
    apply_title_slide,
)


class PptxExporter:
    """Export a LessonPackage to a polished, editable .pptx file."""

    def __init__(self) -> None:
        self._presentation: Presentation | None = None

    def export(self, package: LessonPackage, output_path: Path | str) -> Path:
        out = Path(output_path)
        prs = Presentation()
        # Clean 16:9 canvas.
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT

        blank_layout = prs.slide_layouts[6]  # blank
        ref_lookup = {r.citation_id: r for r in package.references}

        for idx, slide in enumerate(package.slides):
            s = prs.slides.add_slide(blank_layout)
            slide_num = slide.slide_number
            layout = self._classify_slide(slide)
            footer = self._build_footer(slide.citation_ids, ref_lookup)

            if layout == "title":
                apply_title_slide(
                    s,
                    title=slide.title,
                    subtitle=package.metadata.audience,
                )
            elif layout == "objectives":
                apply_objectives_slide(
                    s,
                    title=slide.title,
                    bullets=slide.key_points,
                    footer=footer,
                    slide_number=slide_num,
                )
            elif layout == "evidence":
                evidence_strength, limitation = self._build_evidence_cues(
                    slide.citation_ids, ref_lookup
                )
                apply_evidence_slide(
                    s,
                    title=slide.title,
                    bullets=slide.key_points,
                    footer=footer,
                    evidence_strength=evidence_strength,
                    limitation=limitation,
                    slide_number=slide_num,
                )
            elif layout == "summary":
                apply_summary_slide(
                    s,
                    title=slide.title,
                    bullets=slide.key_points,
                    footer=footer,
                    slide_number=slide_num,
                )
            elif layout == "references":
                apply_references_slide(
                    s,
                    title=slide.title,
                    bullets=slide.key_points,
                    slide_number=slide_num,
                )
            else:
                apply_content_slide(
                    s,
                    title=slide.title,
                    bullets=slide.key_points,
                    footer=footer,
                    slide_number=slide_num,
                )

            # Add speaker notes if python-pptx supports them.
            if slide.speaker_notes:
                try:
                    notes_slide = s.notes_slide
                    notes_frame = notes_slide.notes_text_frame
                    notes_frame.text = slide.speaker_notes
                except Exception:
                    pass  # Notes are also exported to speaker_notes.md.

        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        return out

    @staticmethod
    def _classify_slide(slide: LessonSlide) -> str:
        """Route a slide to a visual layout without changing the schema."""
        title_lower = slide.title.lower()
        slide_type = slide.slide_type.lower()

        if slide_type == "title" or slide.slide_number == 1:
            return "title"
        if slide_type == "references" or "reference" in title_lower:
            return "references"
        if slide_type == "summary" or "summary" in title_lower:
            return "summary"
        if (
            slide_type == "evidence"
            or ("evidence" in title_lower and "limitation" in title_lower)
        ):
            return "evidence"
        if "learning objective" in title_lower or "objectives" in title_lower:
            return "objectives"
        return "content"

    @staticmethod
    def _build_footer(
        citation_ids: list[str],
        ref_lookup: dict[str, CitationRecord],
    ) -> str:
        if not citation_ids:
            return ""
        short_refs = []
        for cid in citation_ids:
            ref = ref_lookup.get(cid)
            if ref:
                if ref.pmid:
                    short_refs.append(f"PMID: {ref.pmid}")
                elif ref.year:
                    short_refs.append(f"{cid} {ref.year}")
                else:
                    short_refs.append(cid)
            else:
                short_refs.append(cid)
        return "Sources: " + "; ".join(short_refs)

    @staticmethod
    def _build_evidence_cues(
        citation_ids: list[str],
        ref_lookup: dict[str, CitationRecord],
    ) -> tuple[str, str]:
        """Derive conservative display cues from existing citation metadata only."""
        cited = [ref_lookup[cid] for cid in citation_ids if cid in ref_lookup]
        if not cited:
            return "Not classified", "Not stated in citation metadata."

        levels = [
            ref.evidence_level.value
            for ref in cited
            if ref.evidence_level
            and ref.evidence_level.value != "unverified / provider-generated"
        ]
        unique_levels = list(dict.fromkeys(levels))
        evidence_strength = "; ".join(unique_levels) or "Not classified"

        limitation = next(
            (ref.limitations for ref in cited if ref.limitations),
            "Not stated in citation metadata.",
        )
        return evidence_strength, limitation
